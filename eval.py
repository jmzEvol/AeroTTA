import os
import os.path as osp
import argparse
from mmengine.runner import Runner
from mmengine.config import Config, DictAction

import segearthov3_segmentor
import custom_datasets

try:
    import openpyxl
except ModuleNotFoundError:
    openpyxl = None


DEFAULT_SEED = 3407


def resolve_eval_work_dir(cfg, config_path):
    """Prefer the LoRA checkpoint directory so eval artifacts sit with training outputs."""
    model_cfg = getattr(cfg, 'model', None)
    lora_path = None
    if model_cfg is not None:
        if isinstance(model_cfg, dict):
            lora_path = model_cfg.get('lora_path')
        else:
            lora_path = getattr(model_cfg, 'lora_path', None)

    if lora_path:
        lora_path = osp.abspath(osp.expanduser(str(lora_path)))
        return osp.dirname(lora_path)

    return osp.join('./work_dirs', osp.splitext(osp.basename(config_path))[0])


def parse_args():
    parser = argparse.ArgumentParser(
        description='CorrCLIP evaluation with MMSeg')
    parser.add_argument('config', default='./configs/cfg_loveda.py')
    parser.add_argument(
        '--show', action='store_true', help='show prediction results')
    parser.add_argument(
        '--show_dir',
        default='./show_dir/',
        help='directory to save visualizaion images')
    parser.add_argument(
        '--out',
        type=str,
        help='The directory to save output prediction for offline evaluation')
    parser.add_argument(
        '--cfg-options',
        nargs='+',
        action=DictAction,
        help='override some settings in the used config, the key-value pair '
        'in xxx=yyy format will be merged into config file. If the value to '
        'be overwritten is a list, it should be like key="[a,b]" or key=a,b '
        'It also allows nested list/tuple values, e.g. key="[(a,b),(c,d)]" '
        'Note that the quotation marks are necessary and that no white space '
        'is allowed.')
    parser.add_argument(
        '--launcher',
        choices=['none', 'pytorch', 'slurm', 'mpi'],
        default='none',
        help='job launcher')
    parser.add_argument(
        '--collect-device',
        choices=['cpu', 'gpu'],
        default='gpu',
        help='device used to collect metrics in distributed evaluation')
    parser.add_argument(
        '--seed',
        type=int,
        default=DEFAULT_SEED,
        help='Base random seed passed to MMEngine for reproducible evaluation.')
    # When using PyTorch version >= 2.0.0, the `torch.distributed.launch`
    # will pass the `--local-rank` parameter to `tools/train.py` instead
    # of `--local_rank`.
    parser.add_argument('--local_rank', '--local-rank', type=int, default=0)
    args = parser.parse_args()
    if 'LOCAL_RANK' not in os.environ:
        os.environ['LOCAL_RANK'] = str(args.local_rank)

    return args


def append_experiment_result(file_path, experiment_data):
    if openpyxl is None:
        return

    try:
        workbook = openpyxl.load_workbook(file_path)
    except FileNotFoundError:
        workbook = openpyxl.Workbook()

    sheet = workbook.active

    if sheet['A1'].value is None:
        sheet['A1'] = 'Model'
        sheet['B1'] = 'Dataset'
        sheet['C1'] = 'aAcc'
        sheet['D1'] = 'mIoU'
        sheet['E1'] = 'mAcc'

    last_row = sheet.max_row

    for index, result in enumerate(experiment_data, start=1):
        sheet.cell(row=last_row + index, column=1, value=result['Model'])
        sheet.cell(row=last_row + index, column=2, value=result['Dataset'])
        sheet.cell(row=last_row + index, column=3, value=result['aAcc'])
        sheet.cell(row=last_row + index, column=4, value=result['mIoU'])
        sheet.cell(row=last_row + index, column=5, value=result['mAcc'])

    workbook.save(file_path)


def trigger_visualization_hook(cfg, args):
    default_hooks = cfg.default_hooks
    if 'visualization' in default_hooks:
        visualization_hook = default_hooks['visualization']
        # Turn on visualization
        visualization_hook['draw'] = True
        if args.show:
            visualization_hook['show'] = True
            visualization_hook['wait_time'] = args.wait_time
        if args.show_dir:
            visualizer = cfg.visualizer
            visualizer['save_dir'] = args.show_dir
    else:
        raise RuntimeError(
            'VisualizationHook must be included in default_hooks.'
            'refer to usage '
            '"visualization=dict(type=\'VisualizationHook\')"')

    return cfg


def main():
    args = parse_args()
    print(os.getcwd())
    cfg = Config.fromfile(args.config)
    cfg.launcher = args.launcher
    # add output_dir in metric
    if args.out is not None:
        cfg.test_evaluator['output_dir'] = args.out
        cfg.test_evaluator['keep_results'] = True
    if args.cfg_options is not None:
        cfg.merge_from_dict(args.cfg_options)
    cfg.work_dir = resolve_eval_work_dir(cfg, args.config)
    os.makedirs(cfg.work_dir, exist_ok=True)
    cfg.randomness = dict(seed=args.seed)

    cfg.test_evaluator['collect_device'] = args.collect_device

    # trigger_visualization_hook(cfg, args)
    runner = Runner.from_cfg(cfg)
    results = runner.test()

    results.update({'Model': cfg.model.model_type,
                    'Dataset': cfg.dataset_type})

    if runner.rank == 0:
        append_experiment_result(osp.join(cfg.work_dir, 'results.xlsx'), [results])

    if runner.rank == 0:
        with open(os.path.join(cfg.work_dir, 'results.txt'), 'a') as f:
            f.write(os.path.basename(args.config).split('.')[0] + '\n')
            for k, v in results.items():
                f.write(k + ': ' + str(v) + '\n')


if __name__ == '__main__':
    main()
